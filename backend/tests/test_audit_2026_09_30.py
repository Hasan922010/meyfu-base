"""Audit №2 (2026-09-30) regressiya testlari — docs/AUDIT_REPORT_2026-09-30.md."""
import uuid
from decimal import Decimal

import pytest

from apps.sales.models import Sale
from apps.warehouse.models import VanStock

SYNC_URL = "/api/v1/sales/bulk-sync/"


def _sync_sale(api, client, product, **item):
    op = {
        "type": "sale",
        "client_uuid": str(uuid.uuid4()),
        "payload": {
            "client": str(client.id),
            "payment_type": "NAQD",
            "items": [{"product": str(product.id), **item}],
        },
    }
    resp = api.post(SYNC_URL, {"operations": [op]}, format="json")
    assert resp.status_code == 200
    return resp.data["data"]["results"][0]


def _van_qty(distributor, product) -> Decimal:
    return VanStock.objects.get(distributor=distributor, product=product).quantity


# ---------- SEC-101: sync orqali manfiy miqdor ----------

@pytest.mark.django_db
def test_sync_sale_negative_quantity_rejected_and_van_unchanged(auth_api, van_stocked):
    product, client = van_stocked["product"], van_stocked["client"]
    before = _van_qty(van_stocked["distributor"], product)

    result = _sync_sale(auth_api, client, product, quantity="-50", price="27000")

    assert result["status"] == "FAILED"
    assert result["error"]["code"] == "INVALID_QUANTITY"
    assert _van_qty(van_stocked["distributor"], product) == before
    assert not Sale.objects.exists()


@pytest.mark.django_db
def test_sync_sale_negative_price_rejected(auth_api, van_stocked):
    result = _sync_sale(auth_api, van_stocked["client"], van_stocked["product"],
                        quantity="1", price="-1000")
    assert result["error"]["code"] == "INVALID_PRICE"


@pytest.mark.django_db
def test_sync_expense_negative_amount_rejected(auth_api, van_stocked, expense_categories):
    op = {
        "type": "expense", "client_uuid": str(uuid.uuid4()),
        "payload": {"category": str(expense_categories["fuel"].id), "amount": "-45000",
                    "payment_source": "CASH_ON_HAND", "description": "x"},
    }
    resp = auth_api.post(SYNC_URL, {"operations": [op]}, format="json")
    assert resp.data["data"]["results"][0]["error"]["code"] == "INVALID_AMOUNT"


# ---------- SEC-102: chegirma bilan min narxni aylanib o'tish ----------

@pytest.mark.django_db
def test_discount_percent_over_100_rejected(auth_api, van_stocked):
    result = _sync_sale(auth_api, van_stocked["client"], van_stocked["product"],
                        quantity="1", price="27000", discount_percent="150")
    assert result["error"]["code"] == "INVALID_DISCOUNT"


@pytest.mark.django_db
def test_discount_below_min_price_rejected_online(auth_api, van_stocked):
    # min_price=24000; 27000 × (1 − 50%) = 13500 < 24000
    resp = auth_api.post("/api/v1/sales/", {
        "client": str(van_stocked["client"].id), "payment_type": "NAQD",
        "items": [{"product": str(van_stocked["product"].id), "quantity": "1",
                   "price": "27000", "discount_percent": "50"}],
    }, format="json")
    assert resp.status_code == 409
    assert resp.data["error"]["code"] == "PRICE_BELOW_MINIMUM"


@pytest.mark.django_db
def test_sale_level_discount_below_min_rejected(auth_api, van_stocked):
    resp = auth_api.post("/api/v1/sales/", {
        "client": str(van_stocked["client"].id), "payment_type": "NAQD",
        "items": [{"product": str(van_stocked["product"].id), "quantity": "2",
                   "price": "27000"}],
        "discount_amount": "50000",
    }, format="json")
    assert resp.status_code == 409
    assert resp.data["error"]["code"] == "PRICE_BELOW_MINIMUM"


@pytest.mark.django_db
def test_serializer_rejects_discount_over_100(auth_api, van_stocked):
    resp = auth_api.post("/api/v1/sales/", {
        "client": str(van_stocked["client"].id), "payment_type": "NAQD",
        "items": [{"product": str(van_stocked["product"].id), "quantity": "1",
                   "price": "27000", "discount_percent": "101"}],
    }, format="json")
    assert resp.status_code == 400


# ---------- BE-115: bir mahsulot ikki qatorda ----------

@pytest.mark.django_db
def test_duplicate_lines_checked_against_total_quantity(auth_api, van_stocked):
    product = van_stocked["product"]
    resp = auth_api.post("/api/v1/sales/", {
        "client": str(van_stocked["client"].id), "payment_type": "NAQD",
        "items": [
            {"product": str(product.id), "quantity": "300", "price": "27000"},
            {"product": str(product.id), "quantity": "300", "price": "27000"},
        ],
    }, format="json")
    assert resp.status_code == 409
    assert resp.data["error"]["details"]["requested"] == "600.000"


# ---------- SEC-103: Telegram bog'lash kodini taxmin qilish ----------

def _tg_msg(chat_id: int, text: str) -> dict:
    return {"message": {"chat": {"id": chat_id}, "message_id": 1, "text": text}}


@pytest.mark.django_db
def test_telegram_link_bruteforce_is_throttled(distributor):
    from apps.telegram_bot.models import TelegramLinkCode
    from apps.telegram_bot.services.webhook import (
        LINK_MAX_FAILS_PER_CHAT,
        handle_update,
    )

    for _ in range(LINK_MAX_FAILS_PER_CHAT):
        handle_update(_tg_msg(777001, "000000"))
    code = TelegramLinkCode.issue(distributor)
    handle_update(_tg_msg(777001, code.code))  # to'g'ri kod ham — bloklangan

    distributor.refresh_from_db()
    assert distributor.telegram_chat_id == ""


@pytest.mark.django_db
def test_telegram_relink_notifies_previous_chat(distributor, monkeypatch):
    from apps.telegram_bot.models import TelegramLinkCode
    from apps.telegram_bot.services import webhook

    sent: list[tuple] = []
    monkeypatch.setattr(
        webhook, "tg_send_chat", lambda chat, text: sent.append((chat, text))
    )
    distributor.telegram_chat_id = "111"
    distributor.save()

    code = TelegramLinkCode.issue(distributor)
    webhook.handle_update(_tg_msg(222, f"/start {code.code}"))

    assert any(chat == "111" for chat, _ in sent)


# ---------- SEC-105: standart webhook siri ----------

@pytest.mark.django_db
def test_webhook_rejects_default_secret(api, settings):
    settings.TELEGRAM_WEBHOOK_SECRET = "dev-webhook-secret"
    resp = api.post(
        "/api/v1/telegram/webhook/", _tg_msg(1, "/help"), format="json",
        HTTP_X_TELEGRAM_BOT_API_SECRET_TOKEN="dev-webhook-secret",
    )
    assert resp.status_code == 403


# ---------- SEC-121: filialsiz filial rahbari botda ----------

@pytest.mark.django_db
def test_branchless_branch_manager_denied_in_bot(monkeypatch):
    from apps.telegram_bot.services import webhook
    from apps.users.models import User

    User.objects.create_user(
        phone="+998907770011", password="pass12345", full_name="BM",
        role="BRANCH_MANAGER", telegram_chat_id="4242",
    )
    sent: list[str] = []
    monkeypatch.setattr(webhook, "tg_send_chat", lambda chat, text: sent.append(text))

    webhook.handle_update(_tg_msg(4242, "/hisobot"))

    assert sent and "faqat admin" in sent[-1]


# ---------- BE-103: kompaniya xarajati kassa jurnalidan ajralmasin ----------

@pytest.mark.django_db
def test_company_expense_cannot_be_patched_or_deleted(admin_api):
    created = admin_api.post("/api/v1/company-expenses/", {
        "category": "RENT", "amount": "1000000", "description": "Ijara",
        "paid_from_cash": True,
    }, format="json")
    assert created.status_code == 201, created.data
    url = f"/api/v1/company-expenses/{created.data['data']['id']}/"

    assert admin_api.patch(url, {"amount": "100"}, format="json").status_code == 405
    assert admin_api.delete(url).status_code == 405


# ---------- BE-102: kechki qaytarish mashina qoldig'idan oshmasin ----------

@pytest.mark.django_db
def test_day_close_return_more_than_van_rejected(auth_api, van_stocked):
    from apps.dayclose.models import DayClose
    from apps.warehouse.models import Stock

    warehouse, product = van_stocked["warehouse"], van_stocked["product"]
    stock_before = Stock.objects.get(warehouse=warehouse, product=product).quantity

    resp = auth_api.post("/api/v1/day-close/submit/", {
        "warehouse": str(warehouse.id), "cash_handed": "0",
        "items": [{"product": str(product.id), "quantity": "501", "condition": "GOOD"}],
    }, format="json")

    assert resp.status_code == 409
    assert resp.data["error"]["code"] == "INSUFFICIENT_STOCK"
    assert not DayClose.objects.exists()
    assert Stock.objects.get(warehouse=warehouse, product=product).quantity == stock_before


@pytest.mark.django_db
def test_day_close_return_to_other_branch_warehouse_rejected(auth_api, van_stocked):
    from apps.warehouse.models import Warehouse

    other = Warehouse.objects.create(name="Begona filial", is_branch=True)
    resp = auth_api.post("/api/v1/day-close/submit/", {
        "warehouse": str(other.id), "cash_handed": "0",
        "items": [{"product": str(van_stocked["product"].id), "quantity": "1",
                   "condition": "GOOD"}],
    }, format="json")

    assert resp.status_code == 409
    assert resp.data["error"]["code"] == "WRONG_WAREHOUSE"


# ---------- BE-101: yopilgan kun qulfi (CLAUDE.md 7.7) ----------

def _close_day(distributor, date):
    from apps.dayclose.constants import DayCloseStatus
    from apps.dayclose.models import DayClose

    return DayClose.objects.create(
        distributor=distributor, date=date, status=DayCloseStatus.CLOSED,
    )


def _sale_body(van_stocked, **extra):
    return {
        "client": str(van_stocked["client"].id), "payment_type": "NAQD",
        "items": [{"product": str(van_stocked["product"].id), "quantity": "1",
                   "price": "27000"}],
        **extra,
    }


@pytest.mark.django_db
def test_online_sale_on_closed_day_rejected(auth_api, van_stocked):
    from apps.core.business_day import business_date

    _close_day(van_stocked["distributor"], business_date())

    resp = auth_api.post("/api/v1/sales/", _sale_body(van_stocked), format="json")

    assert resp.status_code == 409
    assert resp.data["error"]["code"] == "DAY_CLOSED"


@pytest.mark.django_db
def test_offline_sale_for_closed_day_moves_to_today_and_flags(auth_api, van_stocked):
    import datetime

    from apps.core.business_day import business_date
    from apps.core.models import AuditLog

    today = business_date()
    yesterday = today - datetime.timedelta(days=1)
    _close_day(van_stocked["distributor"], yesterday)

    result = _sync_sale(auth_api, van_stocked["client"], van_stocked["product"],
                        quantity="1", price="27000")
    assert result["status"] == "SENT"  # sanasiz — bugun, qulf yo'q
    op = {
        "type": "sale", "client_uuid": str(uuid.uuid4()),
        "payload": {**_sale_body(van_stocked), "date": str(yesterday)},
    }
    res = auth_api.post(SYNC_URL, {"operations": [op]}, format="json")

    assert res.data["data"]["results"][0]["status"] == "SENT"
    sale = Sale.objects.get(client_uuid=op["client_uuid"])
    assert sale.date == today
    assert "DAY_CLOSED" in sale.flag_reason
    assert AuditLog.objects.filter(action="dayclose.late_operation").exists()


@pytest.mark.django_db
def test_cancel_sale_on_closed_day_requires_super_admin_reason(
    auth_api, manager_api, admin_api, van_stocked
):
    from apps.core.business_day import business_date

    created = auth_api.post("/api/v1/sales/", _sale_body(van_stocked), format="json")
    sale_id = created.data["data"]["id"]
    _close_day(van_stocked["distributor"], business_date())
    url = f"/api/v1/sales/{sale_id}/cancel/"

    assert manager_api.post(url, {"reason": "xato"}, format="json").status_code == 409
    assert admin_api.post(url, {}, format="json").status_code == 409
    assert admin_api.post(url, {"reason": "Mijoz qaytardi"},
                          format="json").status_code == 200


# ---------- SEC-118 / BE-113 / SEC-120 ----------

def test_xlsx_export_neutralises_formulas():
    import io

    from openpyxl import load_workbook

    from apps.reports.export import rows_to_xlsx

    data = rows_to_xlsx([["Mijoz", "Summa"], ['=HYPERLINK("http://x","ok")', 5000]])
    ws = load_workbook(io.BytesIO(data)).active

    assert ws["A2"].value.startswith("'=")
    assert ws["A2"].data_type == "s"
    assert ws["B2"].value == 5000


@pytest.mark.django_db
def test_report_range_is_capped(manager_api):
    resp = manager_api.get(
        "/api/v1/reports/export-1c/?date_from=2020-01-01&date_to=2026-01-01"
    )
    assert resp.status_code == 400

    reversed_range = manager_api.get(
        "/api/v1/reports/export-1c/?date_from=2026-02-01&date_to=2026-01-01"
    )
    assert reversed_range.status_code == 400


@pytest.mark.django_db
def test_change_password_rejects_weak_and_revokes_sessions(auth_api, distributor):
    from rest_framework_simplejwt.token_blacklist.models import OutstandingToken

    weak = auth_api.post("/api/v1/auth/change-password/", {
        "old_password": "pass12345", "new_password": "12345678",
    }, format="json")
    assert weak.status_code == 400

    ok_resp = auth_api.post("/api/v1/auth/change-password/", {
        "old_password": "pass12345", "new_password": "Kuchli.Parol-2026",
    }, format="json")
    assert ok_resp.status_code == 200, ok_resp.data
    tokens = OutstandingToken.objects.filter(user=distributor)
    assert tokens.exists()
    assert all(hasattr(t, "blacklistedtoken") for t in tokens)


# ---------- BE-104 / BE-112: audit va ombor yopish ----------

@pytest.mark.django_db
def test_price_change_writes_audit_and_history(admin_api, catalog):
    from apps.catalog.models import ProductPrice
    from apps.core.models import AuditLog

    product = catalog["product"]
    resp = admin_api.patch(f"/api/v1/products/{product.pk}/",
                           {"retail_price": "29000"}, format="json")

    assert resp.status_code == 200, resp.data
    log = AuditLog.objects.get(action="product.price_changed")
    assert log.changes["retail_price"] == ["28000.00", "29000.00"]
    assert ProductPrice.objects.filter(product=product, reason="Qo'lda o'zgartirildi").exists()


@pytest.mark.django_db
def test_client_block_is_audited(manager_api, routed_clients):
    from apps.core.models import AuditLog

    client = routed_clients["my_client"]
    resp = manager_api.patch(f"/api/v1/clients/{client.pk}/",
                             {"is_blocked": True}, format="json")

    assert resp.status_code == 200, resp.data
    assert AuditLog.objects.filter(action="client.updated",
                                   object_id=str(client.pk)).exists()


@pytest.mark.django_db
def test_warehouse_with_stock_cannot_be_deleted(admin_api, stocked):
    resp = admin_api.delete(f"/api/v1/warehouses/{stocked['warehouse'].pk}/")
    assert resp.status_code == 409
    assert resp.data["error"]["code"] == "WAREHOUSE_NOT_EMPTY"
