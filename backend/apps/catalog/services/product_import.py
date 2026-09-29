"""Mahsulotlarni Excel'dan import qilish (v5: B1).

Qoidalar:
  - Kalit — SKU (artikul): bor bo'lsa yangilanadi, yo'q bo'lsa yaratiladi.
  - Kategoriya / birlik / brend nomi bo'yicha topiladi, yo'q bo'lsa yaratiladi.
  - Mavjud mahsulot narxini faqat SUPER_ADMIN o'zgartiradi (CLAUDE.md 2) —
    boshqalarda narx farq qilsa, qator xato bo'ladi.
  - Narx o'zgarsa — ProductPrice tarixi + AuditLog (CLAUDE.md 5.3).
  - `dry_run=True` — hech narsa yozilmaydi, faqat natija ko'rsatiladi.
  - Butun fayl bitta tranzaksiyada: xatoli qator bo'lsa, hech narsa yozilmaydi.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import IO, Any

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.exceptions import BusinessError
from apps.core.models import AuditLog
from apps.users.constants import Role

from ..models import Brand, Category, Product, ProductPrice, Unit

MAX_ROWS = 5000
MAX_FILE_BYTES = 5 * 1024 * 1024
_ZERO = Decimal("0")

# Ustun sarlavhasi (kichik harfda) → model maydoni. O'zbekcha va inglizcha nomlar.
COLUMNS: dict[str, str] = {
    "sku": "sku", "artikul": "sku",
    "name": "name", "nomi": "name",
    "category": "category", "kategoriya": "category",
    "unit": "unit", "birlik": "unit",
    "brand": "brand", "brend": "brand",
    "barcode": "barcode", "shtrix-kod": "barcode", "shtrixkod": "barcode",
    "cost_price": "cost_price", "tannarx": "cost_price",
    "wholesale_price": "wholesale_price", "optom narx": "wholesale_price",
    "retail_price": "retail_price", "chakana narx": "retail_price",
    "min_price": "min_price", "minimal narx": "min_price",
    "pack_quantity": "pack_quantity", "qadoqdagi soni": "pack_quantity",
}
TEMPLATE_HEADER = [
    "artikul", "nomi", "kategoriya", "birlik", "brend", "shtrix-kod",
    "tannarx", "optom narx", "chakana narx", "minimal narx", "qadoqdagi soni",
]
PRICE_COLUMNS = ("cost_price", "wholesale_price", "retail_price", "min_price")
_NUMERIC = (*PRICE_COLUMNS, "pack_quantity")


@dataclass
class RowResult:
    row: int
    sku: str
    action: str  # CREATE | UPDATE | SKIP | ERROR
    errors: list[str] = field(default_factory=list)


@dataclass
class ImportReport:
    dry_run: bool
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    rows: list[RowResult] = field(default_factory=list)

    @property
    def error_count(self) -> int:
        return sum(1 for r in self.rows if r.action == "ERROR")

    def as_dict(self) -> dict[str, Any]:
        return {
            "dry_run": self.dry_run,
            "created": self.created,
            "updated": self.updated,
            "unchanged": self.unchanged,
            "errors": self.error_count,
            "rows": [r.__dict__ for r in self.rows],
        }


class _Rollback(Exception):
    """dry_run yoki xato bo'lganda tranzaksiyani bekor qilish uchun."""


def _read_rows(file: IO[bytes]) -> list[tuple[int, dict[str, Any]]]:
    if getattr(file, "size", 0) > MAX_FILE_BYTES:
        raise BusinessError(message="Fayl hajmi 5 MB dan oshmasin.", code="FILE_TOO_LARGE")
    try:
        wb = load_workbook(file, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl turli xatolar tashlaydi
        raise BusinessError(
            message="Faylni o'qib bo'lmadi. .xlsx formatidagi Excel fayl yuklang.",
            code="INVALID_FILE",
        ) from exc
    ws = wb.active
    iterator = ws.iter_rows(values_only=True)
    header = next(iterator, None)
    if not header:
        raise BusinessError(message="Fayl bo'sh.", code="EMPTY_FILE")
    mapping = {
        idx: COLUMNS[str(h).strip().lower()]
        for idx, h in enumerate(header)
        if h is not None and str(h).strip().lower() in COLUMNS
    }
    if not {"sku", "name"} <= set(mapping.values()):
        raise BusinessError(
            message="Sarlavhada kamida «artikul» va «nomi» ustunlari bo'lishi kerak.",
            code="MISSING_COLUMNS",
        )
    rows: list[tuple[int, dict[str, Any]]] = []
    for number, values in enumerate(iterator, start=2):
        if values is None or all(v in (None, "") for v in values):
            continue
        rows.append((number, {
            mapping[i]: v for i, v in enumerate(values) if i in mapping
        }))
        if len(rows) > MAX_ROWS:
            raise BusinessError(
                message=f"Bir faylda ko'pi bilan {MAX_ROWS} qator bo'lsin.",
                code="TOO_MANY_ROWS",
            )
    wb.close()
    return rows


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel SKU/shtrix-kodni son qilib saqlaydi
    return str(value).strip()


def _parse(raw: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    data: dict[str, Any] = {k: _text(raw.get(k)) for k in
                            ("sku", "name", "category", "unit", "brand", "barcode")}
    errors: list[str] = []
    if not data["sku"]:
        errors.append("Artikul bo'sh.")
    if not data["name"]:
        errors.append("Nomi bo'sh.")
    for key in _NUMERIC:
        text = _text(raw.get(key)).replace(" ", "").replace(",", ".")
        if text == "":
            continue
        try:
            number = Decimal(text)
        except InvalidOperation:
            errors.append(f"«{key}» son emas: {text}")
            continue
        if number < _ZERO or (key == "pack_quantity" and number <= _ZERO):
            errors.append(f"«{key}» manfiy yoki nol bo'lmasin.")
            continue
        data[key] = number
    return data, errors


class _Lookups:
    """Nom bo'yicha kategoriya/birlik/brendni keshlab topadi yoki yaratadi."""

    def __init__(self) -> None:
        self._cache: dict[tuple[str, str], Any] = {}

    def get(self, model: type, name: str) -> Any:
        key = (model.__name__, name.lower())
        if key not in self._cache:
            # o'chirilganini ham qidiramiz — unique nom bilan to'qnashmaslik uchun
            obj = model.all_objects.filter(name__iexact=name).first()
            if obj is not None and obj.is_deleted:
                obj.is_deleted = False
                obj.save(update_fields=["is_deleted", "updated_at"])
            if obj is None:
                extra = {"short_name": name[:16]} if model is Unit else {}
                obj = model.objects.create(name=name, **extra)
            self._cache[key] = obj
        return self._cache[key]


def _apply_row(
    data: dict[str, Any], lookups: _Lookups, user, can_price: bool,
) -> tuple[str, list[str], dict[str, Any] | None]:
    """Bitta qatorni yozadi. (action, errors, narx o'zgarishi) qaytaradi."""
    # SKU unique — o'chirilgan mahsulot qayta import qilinsa, tiklanadi
    product = Product.all_objects.filter(sku=data["sku"]).first()
    fields: dict[str, Any] = {"name": data["name"], "is_deleted": False}
    if data["barcode"]:
        fields["barcode"] = data["barcode"]
    if data["category"]:
        fields["category"] = lookups.get(Category, data["category"])
    if data["unit"]:
        fields["unit"] = lookups.get(Unit, data["unit"])
    if data["brand"]:
        fields["brand"] = lookups.get(Brand, data["brand"])
    for key in _NUMERIC:
        if key in data:
            fields[key] = data[key]

    if product is None:
        missing = [label for key, label in (("category", "kategoriya"), ("unit", "birlik"))
                   if key not in fields]
        if missing:
            return "ERROR", [f"Yangi mahsulot uchun {', '.join(missing)} kerak."], None
        errors = _min_price_errors(fields, None)
        if errors:
            return "ERROR", errors, None
        fields.pop("is_deleted")
        Product.objects.create(sku=data["sku"], created_by=user, **fields)
        return "CREATE", [], None

    changed = {k: v for k, v in fields.items() if getattr(product, k) != v}
    price_changed = [k for k in PRICE_COLUMNS if k in changed]
    if price_changed and not can_price:
        return "ERROR", ["Mavjud mahsulot narxini faqat SUPER_ADMIN o'zgartira oladi."], None
    if not changed:
        return "SKIP", [], None
    errors = _min_price_errors(fields, product)
    if errors:
        return "ERROR", errors, None
    before = {k: str(getattr(product, k)) for k in price_changed}
    for key, value in changed.items():
        setattr(product, key, value)
    product.save()
    if not price_changed:
        return "UPDATE", [], None
    ProductPrice.objects.create(
        product=product, effective_from=timezone.now(), reason="Excel import",
        **{k: getattr(product, k) for k in PRICE_COLUMNS},
    )
    return "UPDATE", [], {
        "sku": product.sku, "before": before,
        "after": {k: str(getattr(product, k)) for k in price_changed},
    }


def _min_price_errors(fields: dict[str, Any], product: Product | None) -> list[str]:
    retail = fields.get("retail_price", getattr(product, "retail_price", None))
    minimum = fields.get("min_price", getattr(product, "min_price", None))
    if retail is not None and minimum is not None and retail > _ZERO and minimum > retail:
        return ["Minimal narx chakana narxdan katta bo'lmasin."]
    return []


def import_products(file: IO[bytes], *, user, dry_run: bool) -> ImportReport:
    rows = _read_rows(file)
    report = ImportReport(dry_run=dry_run)
    can_price = user.is_superuser or getattr(user, "role", None) == Role.SUPER_ADMIN
    price_changes: list[dict[str, Any]] = []
    seen: set[str] = set()

    try:
        with transaction.atomic():
            lookups = _Lookups()
            for number, raw in rows:
                data, errors = _parse(raw)
                if data["sku"] and data["sku"] in seen:
                    errors.append("Bu artikul faylda takrorlangan.")
                seen.add(data["sku"])
                if errors:
                    report.rows.append(RowResult(number, data["sku"], "ERROR", errors))
                    continue
                action, errors, change = _apply_row(data, lookups, user, can_price)
                report.rows.append(RowResult(number, data["sku"], action, errors))
                if change:
                    price_changes.append(change)
            report.created = sum(1 for r in report.rows if r.action == "CREATE")
            report.updated = sum(1 for r in report.rows if r.action == "UPDATE")
            report.unchanged = sum(1 for r in report.rows if r.action == "SKIP")
            if dry_run or report.error_count:
                raise _Rollback
            AuditLog.objects.create(
                user=user, action="product.import", model_name="Product",
                changes={
                    "created": report.created, "updated": report.updated,
                    "price_changes": price_changes[:200],
                },
            )
    except _Rollback:
        pass
    return report
