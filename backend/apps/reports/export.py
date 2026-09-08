"""Excel eksport (openpyxl) — CLAUDE.md 10, 20."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def rows_to_xlsx(rows: list[list], *, sheet_name: str = "Hisobot") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            if r_idx == 1:
                cell.font = Font(bold=True)

    # ustun kengligi
    if rows:
        for c_idx in range(1, len(rows[0]) + 1):
            width = max(
                (len(str(row[c_idx - 1])) for row in rows if c_idx - 1 < len(row)),
                default=10,
            )
            ws.column_dimensions[get_column_letter(c_idx)].width = min(width + 3, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
